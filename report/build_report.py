"""Генерація звіту до лабораторної роботи №3 у форматі ДСТУ 3008:2015.

Береться наявний файл з титульною сторінкою, після неї додаються зміст,
вступ, розділи, висновки, перелік джерел і додатки. Усі числові результати
підтягуються з ``data/results.json`` (результат запуску конвеєра), тож
звіт завжди відповідає фактичному запуску.

Зміст будується у два проходи: спершу документ рендериться в PDF
(LibreOffice), за текстом сторінок визначаються номери сторінок заголовків,
потім документ збирається повторно з правильними номерами.

Запуск:  .venv/bin/python report/build_report.py
"""
from __future__ import annotations

import ast
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from docx.shared import Cm, Pt, Twips

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Звіт_Лабораторна_робота_3_Ходаков.docx"
OUT = ROOT / "report" / "Звіт_Лабораторна_робота_3_Ходаков.docx"
FIG = ROOT / "report" / "figures"
SRC = ROOT / "src" / "lab3"
TEXT_WIDTH_DXA = 11906 - 1701 - 850  # ширина тексту між полями шаблону

R = json.loads((ROOT / "data" / "results.json").read_text())
GT = json.loads((ROOT / "data" / "synthetic_scene_ground_truth.json").read_text())
S, EV, META = R["summary"], R["evaluation"], R["meta"]
TRACKS = S["tracks"]
DRV = next(t for t in TRACKS if t["track_id"] == S["driver_id"])
MOV = next(t for t in TRACKS if t["track_id"] in S["moving_ids"])
PAS_SIM = EV["driver"]["similarity"].get("passenger", {"mean": float("nan"), "n": 0})


def f2(x, n=2):
    """Число з десятковою комою (українська нотація)."""
    return f"{x:.{n}f}".replace(".", ",")


# ==========================================================================
# Низькорівневі помічники форматування
# ==========================================================================
class Builder:
    def __init__(self, toc_pages: dict[str, int] | None = None):
        self.doc = Document(str(TEMPLATE))
        self.toc_pages = toc_pages or {}
        self.toc: list[tuple[int, str]] = []  # (рівень, текст)
        self.fig_n: dict[str, int] = {}
        self.tab_n: dict[str, int] = {}
        self.lst_n: dict[str, int] = {}
        self.chapter = "0"
        # Сторінка з титулом закінчується — новий матеріал з нової сторінки.
        # Шаблон уже може закінчуватися розривом сторінки: тоді зайві порожні
        # абзаци після нього прибираємо, а новий розрив не додаємо.
        paras = self.doc.paragraphs
        while paras and not paras[-1].text.strip() and 'w:type="page"' not in paras[-1]._p.xml:
            paras[-1]._p.getparent().remove(paras[-1]._p)
            paras = self.doc.paragraphs
        if 'w:type="page"' not in paras[-1]._p.xml:
            paras[-1].add_run().add_break(WD_BREAK.PAGE)

    # ---------------- абзаци ----------------
    def _runs(self, p, text: str, size: float | None = None):
        """Підтримка розмітки: **жирний**, *курсив*, `код`, _{нижній індекс}."""
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|_\{[^}]+\})", text):
            if not part:
                continue
            if part.startswith("_{"):
                r = p.add_run(part[2:-1])
                r.font.subscript = True
            elif part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = "Courier New"
                r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
                r.font.size = Pt((size or 14) - 1)
            elif part.startswith("*"):
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                r = p.add_run(part)
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
        return p

    def para(self, text: str, align=None, indent=True, size=None, keep_next=False, space_after=0):
        p = self.doc.add_paragraph()
        self._runs(p, text, size)
        pf = p.paragraph_format
        if align is not None:
            p.alignment = align
        if not indent:
            pf.first_line_indent = Cm(0)
        if keep_next:
            pf.keep_with_next = True
        if space_after:
            pf.space_after = Pt(space_after)
        return p

    def bullets(self, items: list[str]):
        """Перелік за ДСТУ 3008: кожен елемент з тире, з абзацного відступу."""
        for i, it in enumerate(items):
            end = "." if i == len(items) - 1 else ";"
            self.para("– " + it.rstrip(".;") + end)

    def numbered(self, items: list[str]):
        for i, it in enumerate(items, 1):
            end = "." if i == len(items) else ";"
            self.para(f"{i}) " + it.rstrip(".;") + end)

    # ---------------- заголовки ----------------
    def _toc_entry(self, level: int, text: str):
        self.toc.append((level, text))

    def structural(self, text: str, toc=True, page_break=True):
        """Заголовок структурного елемента (ВСТУП, ВИСНОВКИ…) — по центру, великими."""
        p = self.doc.add_paragraph(style="Heading 1")
        p.add_run(text.upper())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.page_break_before = page_break
        if toc:
            self._toc_entry(0, text.upper())
        return p

    def h1(self, num: str, text: str):
        """Заголовок розділу: «1 НАЗВА» з абзацного відступу, з нової сторінки."""
        self.chapter = num
        p = self.doc.add_paragraph(style="Heading 1")
        p.add_run(f"{num} {text.upper()}")
        p.paragraph_format.page_break_before = True
        self._toc_entry(0, f"{num} {text.upper()}")

    def h2(self, num: str, text: str):
        p = self.doc.add_paragraph(style="Heading 2")
        p.add_run(f"{num} {text}")
        self._toc_entry(1, f"{num} {text}")

    # ---------------- рисунки, таблиці, формули, лістинги ----------------
    def _next(self, counter: dict) -> str:
        counter[self.chapter] = counter.get(self.chapter, 0) + 1
        return f"{self.chapter}.{counter[self.chapter]}"

    def figure(self, path: Path, caption: str, width_cm: float = 16.0) -> str:
        n = self._next(self.fig_n)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.para(f"Рисунок {n} – {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=8)
        c.paragraph_format.line_spacing = 1.0
        return n

    def next_fig(self) -> str:
        """Номер рисунка, який буде вставлено наступним (для посилань у тексті)."""
        return f"{self.chapter}.{self.fig_n.get(self.chapter, 0) + 1}"

    def next_tab(self) -> str:
        return f"{self.chapter}.{self.tab_n.get(self.chapter, 0) + 1}"

    def next_lst(self) -> str:
        return f"{self.chapter}.{self.lst_n.get(self.chapter, 0) + 1}"

    def table(self, caption: str, header: list[str], rows: list[list], widths_cm: list[float],
              size: float = 12, align_cols: dict | None = None) -> str:
        n = self._next(self.tab_n)
        cap = self.para(f"Таблиця {n} – {caption}", indent=False, keep_next=True)
        cap.paragraph_format.space_before = Pt(6)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t)
        # Колонки масштабуються на всю ширину тексту і фіксуються явно (tblW,
        # tblGrid, fixed layout), інакше Word/LibreOffice підбирають їх за вмістом.
        total = sum(widths_cm)
        _fix_table_widths(t, [int(w / total * TEXT_WIDTH_DXA) for w in widths_cm])
        for ri, row in enumerate([header] + rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                self._runs(p, str(val), size)
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.line_spacing = 1.0
                if ri == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.bold = True
                    _repeat_header(t.rows[0])
                else:
                    p.alignment = (align_cols or {}).get(ci, WD_ALIGN_PARAGRAPH.LEFT)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return n

    def formula(self, omml: str) -> str:
        """Формула (рівняння Word, OMML) по центру, номер «(2.1)» — біля правого поля."""
        key = "_formula_" + self.chapter
        self.tab_n[key] = self.tab_n.get(key, 0) + 1
        n = f"{self.chapter}.{self.tab_n[key]}"
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        tabs = p.paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(16.45 / 2), alignment=1)  # центр
        tabs.add_tab_stop(Cm(16.45), alignment=2)      # праве поле
        p.add_run("\t")
        p._p.append(parse_xml(f'<m:oMath {_NS}>{omml}</m:oMath>'))
        p.add_run(f"\t({n})")
        return n

    def listing(self, caption: str, code: str, size: float = 9.5) -> str:
        n = self._next(self.lst_n)
        cap = self.para(f"Лістинг {n} – {caption}", indent=False, keep_next=True)
        cap.paragraph_format.space_before = Pt(6)
        lines = code.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.0
            pf.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if i < len(lines) - 1 and len(lines) < 40:
                pf.keep_with_next = True
            r = p.add_run(line if line else " ")
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
            r._element.rPr.rFonts.set(qn("w:cs"), "Courier New")
            r.font.size = Pt(size)
            _shade(p, "F2F2F2")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return n

    # ---------------- зміст ----------------
    def glossary(self, items: list[tuple[str, str]]):
        """Перелік скорочень: зліва — скорочення, справа — розшифрування (таблиця без рамок)."""
        t = self.doc.add_table(rows=len(items), cols=2)
        _fix_table_widths(t, [int(TEXT_WIDTH_DXA * 0.22), TEXT_WIDTH_DXA - int(TEXT_WIDTH_DXA * 0.22)])
        for (abbr, text), row in zip(items, t.rows):
            for cell, val in zip(row.cells, (abbr, "– " + text)):
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._runs(p, val)

    def toc_placeholder(self):
        self.structural("Зміст", toc=False, page_break=True)
        self._toc_anchor = self.doc.add_paragraph()

    def fill_toc(self):
        """Вставляє рядки змісту перед якорем (номери сторінок — з першого проходу)."""
        anchor = self._toc_anchor
        for level, text in self.toc:
            p = anchor.insert_paragraph_before()
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0.6 * level)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = p._p.get_or_add_pPr()
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "right")
            tab.set(qn("w:leader"), "dot")
            tab.set(qn("w:pos"), str(TEXT_WIDTH_DXA))
            tabs.append(tab)
            pPr.append(tabs)
            p.add_run(text)
            p.add_run("\t" + str(self.toc_pages.get(text, 0)))
        anchor._p.getparent().remove(anchor._p)

    def save(self, path: Path):
        self.fill_toc()
        _schema_order(self.doc)
        self.doc.save(str(path))


# ---------------- побудова формул OMML ----------------
_NS = ('xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
       'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"')


def mr(text: str, plain: bool = False) -> str:
    """Фрагмент формули; ``plain`` — прямий шрифт (назви функцій, числа)."""
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""
    return (f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
            f'<w:sz w:val="28"/><w:szCs w:val="28"/></w:rPr>'
            f'<m:t xml:space="preserve">{text}</m:t></m:r>')


def msub(base: str, sub: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"


def mfrac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def msqrt(e: str) -> str:
    return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{e}</m:e></m:rad>'


def mdelim(e: str, beg: str = "(", end: str = ")") -> str:
    """Дужки звичайними символами — однаково відображаються у Word і LibreOffice."""
    return mr(beg, True) + e + mr(end, True)


F_SIMILARITY = mr("x′=s") + mr("·") + mr("R") + mdelim(mr("θ")) + mr("·x+t")
F_CHAIN = (msub(mr("T"), mr("t")) + mr("=") + msub(mr("A"), mr("1", True)) + mr("·") +
           msub(mr("A"), mr("2", True)) + mr("·…·") + msub(mr("A"), mr("t")))
F_COSINE = (mr("sim", True) + mdelim(mr("a,b")) + mr("=") +
            mfrac(mr("a·b"), mdelim(mr("a"), "‖", "‖") + mr("·") + mdelim(mr("b"), "‖", "‖")))
F_MOVE = mr("M=") + mfrac(
    msub(mr("median", True), mr("e")) + mdelim(msub(mr("P"), mr("90", True)) + mdelim(mr("e")) + mr("−") +
                                               msub(mr("P"), mr("10", True)) + mdelim(mr("e"))),
    mr("median", True) + mdelim(msqrt(mr("w·h"))))


# Порядок дочірніх елементів, якого вимагає схема OOXML (ECMA-376). Елементи,
# додані через OxmlElement у кінець, переставляються, щоб Word відкривав файл
# без попереджень.
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"]
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
                "tblLook", "tblCaption", "tblDescription"]


def _schema_order(doc):
    body = doc.element.body
    for tag, order in (("w:pPr", _PPR_ORDER), ("w:tblPr", _TBLPR_ORDER)):
        rank = {qn("w:" + n): i for i, n in enumerate(order)}
        for el in body.iter(qn(tag)):
            children = list(el)
            children.sort(key=lambda c: rank.get(c.tag, len(order)))
            for c in children:
                el.remove(c)
            for c in children:
                el.append(c)


def _table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def _fix_table_widths(table, widths_dxa: list[int]):
    """Фіксує ширину таблиці та кожної колонки (у twips/dxa)."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(sum(widths_dxa)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_dxa):
        gc.set(qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_dxa):
            cell.width = Twips(w)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trPr.append(h)


def _shade(p, color: str):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def source_of(module: str, *names: str) -> str:
    """Повертає вихідний код вказаних функцій/класів модуля (для лістингів)."""
    text = (SRC / module).read_text()
    tree = ast.parse(text)
    out = []
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.ClassDef)) and node.name in names:
            out.append(ast.get_source_segment(text, node))
    return "\n\n".join(out)


def fragment(code: str, start: str, end: str | None = None) -> str:
    """Вирізає з коду рядки від першого, що містить ``start``, до рядка з ``end``
    включно; прибирає спільний відступ. Повний код — у додатку А."""
    import textwrap
    lines = code.split("\n")
    i = next(k for k, l in enumerate(lines) if start in l)
    j = next(k for k in range(i, len(lines)) if end in lines[k]) if end else len(lines) - 1
    return textwrap.dedent("\n".join(lines[i:j + 1]))


def method_of(module: str, cls: str, *names: str) -> str:
    text = (SRC / module).read_text()
    tree = ast.parse(text)
    c = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == cls)
    seg = [ast.get_source_segment(text, n, padded=True) for n in c.body
           if isinstance(n, ast.FunctionDef) and n.name in names]
    return "\n\n".join(seg)


REFERENCES = [
    "Jocher G., Qiu J. Ultralytics YOLO11 : software. Version 11.0.0. 2024. URL: "
    "https://github.com/ultralytics/ultralytics (дата звернення: 22.10.2026).",
    "Zhang Y., Sun P., Jiang Y. et al. ByteTrack: Multi-Object Tracking by Associating Every Detection Box. "
    "Computer Vision – ECCV 2022. Cham : Springer, 2022. P. 1–21.",
    "Lin T.-Y., Maire M., Belongie S. et al. Microsoft COCO: Common Objects in Context. Computer Vision – "
    "ECCV 2014. Cham : Springer, 2014. P. 740–755.",
    "morsetechlab. YOLOv11 License Plate Detection : model. Hugging Face, 2024. URL: "
    "https://huggingface.co/morsetechlab/yolov11-license-plate-detection (дата звернення: 22.10.2026).",
    "JaidedAI. EasyOCR : software. Version 1.7.2. URL: https://github.com/JaidedAI/EasyOCR "
    "(дата звернення: 22.10.2026).",
    "Baek Y., Lee B., Han D. et al. Character Region Awareness for Text Detection. Proceedings of the "
    "IEEE/CVF Conference on Computer Vision and Pattern Recognition (CVPR). 2019. P. 9365–9374.",
    "Shi B., Bai X., Yao C. An End-to-End Trainable Neural Network for Image-Based Sequence Recognition and "
    "Its Application to Scene Text Recognition. IEEE Transactions on Pattern Analysis and Machine "
    "Intelligence. 2017. Vol. 39, No. 11. P. 2298–2304.",
    "Wu W., Peng H., Yu S. YuNet: A Tiny Millisecond-level Face Detector. Machine Intelligence Research. "
    "2023. Vol. 20, No. 5. P. 656–665.",
    "Zhong Y., Deng W., Hu J. et al. SFace: Sigmoid-Constrained Hypersphere Loss for Robust Face "
    "Recognition. IEEE Transactions on Image Processing. 2021. Vol. 30. P. 2587–2598.",
    "Lucas B. D., Kanade T. An Iterative Image Registration Technique with an Application to Stereo Vision. "
    "Proceedings of the 7th International Joint Conference on Artificial Intelligence. 1981. P. 674–679.",
    "Shi J., Tomasi C. Good Features to Track. Proceedings of the IEEE Conference on Computer Vision and "
    "Pattern Recognition (CVPR). 1994. P. 593–600.",
    "Fischler M. A., Bolles R. C. Random Sample Consensus: A Paradigm for Model Fitting with Applications "
    "to Image Analysis and Automated Cartography. Communications of the ACM. 1981. Vol. 24, No. 6. "
    "P. 381–395.",
    "Bouguet J.-Y. Pyramidal Implementation of the Lucas Kanade Feature Tracker. Description of the "
    "Algorithm. Intel Corporation, Microprocessor Research Labs, 2001. 9 p.",
    "Bradski G. The OpenCV Library. Dr. Dobb's Journal of Software Tools. 2000. Vol. 25, No. 11. P. 120–125.",
    "OpenCV Zoo: models for OpenCV DNN : repository. URL: https://github.com/opencv/opencv_zoo "
    "(дата звернення: 22.10.2026).",
    "Zuiderveld K. Contrast Limited Adaptive Histogram Equalization. Graphics Gems IV. San Diego : "
    "Academic Press, 1994. P. 474–485.",
    "Демонстрація до завдання № 3. Обробка відео: об'єкти, номер автомобіля, ідентифікація водія : "
    "навчальні матеріали з дисципліни «Методи обробки мультимедійної інформації». Київ, 2026. 6 с.",
    "ДСТУ 4278:2019. Дорожній транспорт. Знаки номерні транспортних засобів. Загальні вимоги. Правила "
    "застосування. Київ : ДП «УкрНДНЦ», 2019.",
    "ДСТУ 3008:2015. Інформація та документація. Звіти у сфері науки і техніки. Структура та правила "
    "оформлювання. Київ : ДП «УкрНДНЦ», 2016. 26 с.",
]

ABBREVIATIONS = [
    ("CLAHE", "Contrast Limited Adaptive Histogram Equalization, адаптивне вирівнювання гістограми з "
              "обмеженням контрасту"),
    ("COCO", "Common Objects in Context, набір даних для виявлення об'єктів (80 класів)"),
    ("CRAFT", "Character Region Awareness for Text detection, детектор тексту"),
    ("CRNN", "Convolutional Recurrent Neural Network, згорткова рекурентна нейронна мережа"),
    ("CTC", "Connectionist Temporal Classification, декодування послідовності без сегментації символів"),
    ("H.264", "стандарт стиснення відео (MPEG-4 AVC)"),
    ("IoU", "Intersection over Union, відношення площі перетину двох рамок до площі їх об'єднання"),
    ("LSTM", "Long Short-Term Memory, рекурентна мережа з довгою короткочасною пам'яттю"),
    ("MPS", "Metal Performance Shaders, обчислення на графічному процесорі Apple Silicon"),
    ("OCR", "Optical Character Recognition, оптичне розпізнавання символів"),
    ("RANSAC", "Random Sample Consensus, метод стійкого оцінювання параметрів моделі"),
    ("SFace", "мережа розпізнавання облич (Sigmoid-Constrained Hypersphere Loss)"),
    ("YOLO", "You Only Look Once, сімейство одностадійних детекторів об'єктів"),
    ("YOLOv11-LP", "YOLOv11, донавчена на номерних знаках (License Plate)"),
    ("YuNet", "легкий згортковий детектор облич з OpenCV Zoo"),
    ("ДСТУ", "національний стандарт України"),
    ("ПЗ", "програмне забезпечення"),
    ("трек", "послідовність положень одного об'єкта в кадрах відео з постійним номером"),
    ("M", "нормована оцінка руху об'єкта після компенсації руху камери"),
    ("sim", "косинусна схожість векторів ознак облич"),
    ("px", "піксель"),
    ("кадр/с", "кадрів за секунду"),
]


# ==========================================================================
# Зміст звіту
# ==========================================================================
def build(toc_pages=None, stats=None) -> Builder:
    b = Builder(toc_pages)
    P = b.para
    stats = stats or {"pages": 0, "figures": 0, "tables": 0}

    # ------------------------------------------------------------------ РЕФЕРАТ
    b.structural("Реферат", toc=False, page_break=False)
    P(f"Звіт про лабораторну роботу: {stats['pages']} с., {stats['figures']} рис., {stats['tables']} табл., "
      f"{len(REFERENCES)} джерел, 2 додатки.")
    P("Об'єкт дослідження — процес обробки відео для виявлення та розпізнавання об'єктів дорожньої сцени.")
    P("Предмет дослідження — методи виявлення та відстеження об'єктів, розпізнавання номерних знаків та "
      "ідентифікації осіб на основі попередньо навчених нейронних мереж.")
    P("Мета — розробити елементи програмного забезпечення, що виділяють і розпізнають об'єкти у відеофільмі, "
      "визначають номер автомобіля, ідентифікують водія та визначають рухомий об'єкт.")
    P("Методи: згорткові нейронні мережі YOLO11m та YOLOv11-LP, трекер ByteTrack, оптичний потік "
      "Лукаса–Канаде з RANSAC, OCR EasyOCR з постобробкою за шаблоном номера та голосуванням по кадрах, "
      "детектор облич YuNet і мережа розпізнавання облич SFace.")
    P(f"Результати: на синтетичному відео тривалістю 40 с з еталонною розміткою виявлено й розпізнано "
      f"{len(TRACKS)} об'єктів сцени, номер автомобіля {S['plate_text']} прочитано точно в "
      f"{f2(EV['plate']['exact_read_rate'] * 100, 0)} % кадрів з відкритим знаком, водія ідентифіковано із "
      f"середньою косинусною схожістю {f2(DRV['mean_similarity'], 3)}, єдиний рухомий об'єкт (пасажир) "
      "визначено завдяки компенсації руху камери.")
    P("ОБРОБКА ВІДЕО, ВИЯВЛЕННЯ ОБ'ЄКТІВ, YOLO11, BYTETRACK, РОЗПІЗНАВАННЯ НОМЕРНИХ ЗНАКІВ, EASYOCR, "
      "РОЗПІЗНАВАННЯ ОБЛИЧ, SFACE, ОПТИЧНИЙ ПОТІК, КОМПЕНСАЦІЯ РУХУ КАМЕРИ")

    b.toc_placeholder()

    # ------------------------------------------------------------------ СКОРОЧЕННЯ
    b.structural("Перелік умовних позначень, символів, одиниць, скорочень і термінів")
    b.glossary(ABBREVIATIONS)

    # ------------------------------------------------------------------ ВСТУП
    b.structural("Вступ")
    P("Обробка відео є однією з найбільш затребуваних задач комп'ютерного зору: системи відеоспостереження, "
      "контролю в'їзду на парковки, автоматичної фіксації порушень правил дорожнього руху та пошуку "
      "транспортних засобів поєднують у собі виявлення об'єктів, їх відстеження між кадрами, читання "
      "номерних знаків та розпізнавання облич. Сучасні попередньо навчені нейронні мережі дозволяють зібрати "
      "такий конвеєр з готових компонентів, а основна інженерна робота полягає в їх правильному поєднанні, "
      "постобробці результатів і прийнятті рішень на рівні всього відео, а не окремого кадру.")
    P("**Мета роботи** — розробити елементи програмного забезпечення для обробки відеофільму тривалістю "
      "30–60 с, що містить автомобіль, його номерний знак, обличчя водія біля автомобіля, додаткові об'єкти та "
      "один рухомий об'єкт, і з використанням відомих нейронних мереж виконати:")
    b.bullets(["виділення об'єктів у фільмі та їх розпізнавання (автомобіль, номерний знак, люди, тварина)",
               "визначення номера автомобіля за допомогою мережі для розпізнавання номерних знаків",
               "ідентифікацію водія"])
    P("Відповідно до дозволу викладача відеофільм є синтетичним: він програмно зібраний з вільних фотографій і "
      "**відрізняється від демонстраційного прикладу** [17] — інший автомобіль (ЗАЗ Sens, вид спереду, на "
      "шкільному подвір'ї), власний номерний знак `KA 7429 MB`, інший водій та інше еталонне фото, пасажир з "
      "іншого джерела, собака та другий автомобіль як додаткові об'єкти, тривалість 40 с замість 10 с, "
      "віртуальна камера «з руки» з наближенням і панорамуванням. Перевагою синтетичного відео є наявність "
      "точної еталонної розмітки для кожного кадру, що дозволило оцінити якість роботи кількісно. Звіт оформлено відповідно "
      "до ДСТУ 3008:2015 [19].")
    P("Для досягнення мети розв'язано такі задачі:")
    b.numbered(["сформовано синтетичне відео 40 с (1280×720, 25 кадр/с) з еталонною розміткою",
                "реалізовано виявлення об'єктів мережею YOLO11m та їх відстеження алгоритмом ByteTrack",
                "реалізовано оцінку руху камери оптичним потоком Лукаса–Канаде з RANSAC для коректного визначення "
                "рухомого об'єкта",
                "реалізовано розпізнавання номерного знака: детектор YOLOv11-LP, OCR EasyOCR, нормалізацію за "
                "шаблоном українського номера та голосування по кадрах",
                "реалізовано ідентифікацію водія мережами YuNet і SFace за окремим еталонним фото",
                "проведено кількісну оцінку результатів, аналіз помилок і запропоновано способи покращення"])

    # ------------------------------------------------------------------ 1
    b.h1("1", "Постановка задачі та вихідні дані")
    b.h2("1.1", "Умова завдання")
    P("Згідно із завданням № 3 з дисципліни «Методи обробки мультимедійної інформації» необхідно зняти "
      "невеликий фільм тривалістю 30–60 с, який містить автомобіль, номер автомобіля, обличчя водія біля "
      "автомобіля та хоча б один додатковий об'єкт (пасажира, другий автомобіль, рюкзак, домашню тварину тощо). "
      "У відео має бути один рухомий об'єкт: машина, людина або тварина. Розроблене ПЗ повинно виділяти вказані "
      "об'єкти та розпізнавати їх, визначати номер автомобіля за допомогою відомих мереж і ідентифікувати водія. "
      "Дозволено використовувати існуючі технології розпізнавання цифр, облич та об'єктів.")
    P("Відповідність вимог завдання вмісту відео наведено в таблиці " + b.next_tab() + ".")
    b.table("Відповідність вмісту відео вимогам завдання",
            ["Вимога завдання", "Реалізація у відео"],
            [["Тривалість 30–60 с", f"{GT['duration_s']} с, {GT['frames']} кадрів, {GT['fps']} кадр/с"],
             ["Автомобіль", "ЗАЗ Sens (вид спереду), займає центр кадру"],
             ["Номер автомобіля", "власний знак KA 7429 MB, ширина в кадрі ≈ 210–290 пікселів"],
             ["Обличчя водія біля автомобіля", "водій стоїть праворуч біля автомобіля, обличчя анфас"],
             ["Додаткові об'єкти", "пасажир, другий (синій) автомобіль, собака"],
             ["Один рухомий об'єкт", "пасажир, що йде зліва направо з 3-ї по 33-тю секунду"],
             ["Окреме фото водія для ідентифікації", "фото того ж водія 2016 р. (інший рік, одяг, освітлення)"]],
            [6.5, 10])

    b.h2("1.2", "Опис відео")
    P("Відео сформоване модулем `scene.py`. Сцена складається з шарів, які накладаються в порядку від "
      "дальнього до ближнього: фон з автомобілями, собака, пасажир, водій. Параметри відео наведено в "
      f"таблиці {b.next_tab()}, джерела зображень — у таблиці {b.next_tab()[:-1]}{int(b.next_tab()[-1]) + 1}.")
    b.table("Параметри відео",
            ["Параметр", "Значення"],
            [["Файл", "data/synthetic_scene.mp4"],
             ["Роздільна здатність", f"{GT['width']}×{GT['height']} пікселів"],
             ["Частота кадрів", f"{GT['fps']} кадр/с"],
             ["Тривалість / кількість кадрів", f"{GT['duration_s']} с / {GT['frames']} кадрів"],
             ["Кодек", "H.264 (libx264), yuv420p, 8 Мбіт/с"],
             ["Рух камери", "наближення 1,00→1,35→1,00; панорама; тремтіння ±7 px; поворот ±0,35°"],
             ["Шумові ефекти", "гаусове розмиття σ=0,6; шум сенсора σ=2,2; коливання експозиції ±2,5 %"]],
            [6, 10.5])
    b.table("Джерела зображень для сцени",
            ["Файл", "Призначення", "Автор, ліцензія"],
            [["car_zaz_sens.jpg", "фон, головний і другий автомобіль", "Вальдимар, CC BY-SA 3.0"],
             ["driver_in_scene_2012.jpg", "водій у відео", "UNICEF Ukraine, CC BY 2.0"],
             ["driver_reference_2016.jpg", "еталонне фото водія", "Oleg Dubyna, CC BY-SA 2.0"],
             ["passenger_source.jpg", "пасажир (рухомий об'єкт)", "Ultralytics, тестове зображення bus.jpg"],
             ["dog_puppy.jpg", "собака (додатковий об'єкт)", "Retyi, CC0"]],
            [5, 6, 5.5])
    P("Головний автомобіль на вихідному фото має зачорнений номерний знак. На його місце нанесено власний "
      "знак `KA 7429 MB`, намальований за пропорціями стандартного українського номерного знака (520×112 мм, "
      "синя смуга з прапором і кодом «UA» зліва) згідно з ДСТУ 4278:2019 [18]. Знак переноситься на фото "
      "перспективним перетворенням (гомографією), що задається чотирма кутами ділянки знака; потім він "
      "затемнюється й трохи розмивається, щоб відповідати освітленню та різкості знімка.")
    P("Водія, пасажира та собаку вирізано з їхніх фотографій сегментаційною мережею YOLO11m-seg: полігон "
      "маски екземпляра заливається в повній роздільній здатності, маска звужується на 2 пікселі та "
      "розмивається, утворюючи м'який альфа-канал. Колірна статистика вирізаних об'єктів частково "
      "наближається до фону переносом середнього та дисперсії в просторі L*a*b*, під пасажира й собаку "
      "додаються м'які тіні. Пасажир рухається з легким вертикальним «підстрибуванням» кроку.")
    P("Віртуальна камера задається ключовими положеннями (центр вікна, наближення) з плавною інтерполяцією "
      "smoothstep, до яких додається тремтіння рук як сума синусоїд різних частот. Для кожного кадру "
      "будується матриця афінного перетворення «світ → кадр», після чого додаються розмиття, шум сенсора та "
      f"коливання експозиції. Приклади кадрів наведено на рисунку {b.next_fig()}.")
    _frames_grid()
    b.figure(FIG / "scene_grid.jpg", "Кадри синтетичного відео: 0 с, 10 с, 18 с і 36 с", 16.2)

    b.h2("1.3", "Еталонна розмітка")
    P("Оскільки сцена формується програмно, для кожного кадру відомі істинні дані, які зберігаються у файлі "
      "`data/synthetic_scene_ground_truth.json`:")
    b.bullets(["рамки видимої частини кожного об'єкта (головний і другий автомобіль, номерний знак, водій, "
               "пасажир, собака) з урахуванням перекриттів шарів і меж кадру",
               "частка видимих пікселів об'єкта (видимість від 0 до 1)",
               "параметри та матриця перетворення камери",
               "текст номера (KA7429MB), який із людей є водієм і який об'єкт рухається"])
    P("Ця розмітка використовується модулем `evaluation.py` для обчислення метрик якості (розділ 4).")

    # ------------------------------------------------------------------ 2
    b.h1("2", "Теоретичні відомості та вибір моделей")
    b.h2("2.1", "Виявлення об'єктів. YOLO11")
    P("YOLO (You Only Look Once) — сімейство одностадійних детекторів, що за один прохід згорткової мережі "
      "передбачають рамки, класи та впевненість для всіх об'єктів зображення [1]. YOLO11 (Ultralytics, 2024) "
      "складається з магістральної мережі (backbone) з блоками C3k2, шиї (neck) типу FPN/PAN з модулем "
      "просторової уваги C2PSA та безякірної (anchor-free) голови, що окремо передбачає класи та координати "
      "рамок на трьох масштабах. Модель навчена на наборі COCO (Common Objects in Context) з 80 класів [3], "
      "серед яких є всі потрібні в завданні: person, car, truck, bus, dog, cat, backpack, handbag тощо.")
    P("Використано варіант YOLO11m (20,1 млн параметрів, mAP 51,5 на COCO val2017) замість YOLO11s з "
      "демонстрації: він точніше розрізняє легковики та вантажівки на незвичних ракурсах і стабільніше "
      "працює з частково перекритими об'єктами, а швидкості на графічному процесорі Apple M1 (MPS) достатньо.")
    b.h2("2.2", "Відстеження об'єктів. ByteTrack")
    P("Детектор працює з кожним кадром незалежно, тож для аналізу поведінки об'єктів детекції потрібно "
      "пов'язати в часі — побудувати треки. ByteTrack [2] прогнозує положення кожного треку фільтром Калмана "
      "і зіставляє прогнози з детекціями за перетином рамок (IoU) угорським алгоритмом. Його особливість — "
      "двоетапне зіставлення: спершу з упевненими детекціями, а потім з детекціями низької впевненості, які "
      "зазвичай відповідають частково перекритим об'єктам. Завдяки цьому об'єкт не «губиться», коли його "
      "ненадовго перекриває інший (у нашому відео пасажир проходить перед собакою та номерним знаком).")
    b.h2("2.3", "Оцінка руху камери")
    P("Під час зйомки «з руки» камера наближається, панорамує і тремтить, тому в координатах кадру "
      "зміщуються всі об'єкти, навіть нерухомі. Щоб відрізнити власний рух об'єкта від руху камери, оцінюється "
      "глобальне перетворення між сусідніми кадрами. На попередньому кадрі детектором Ши–Томасі [11] "
      "знаходяться кутові точки фону (ділянки всередині рамок об'єктів маскуються), їх положення на поточному "
      "кадрі визначається пірамідальним оптичним потоком Лукаса–Канаде [10, 13]. За парами точок методом "
      "RANSAC [12] оцінюється перетворення подібності з чотирма ступенями свободи (масштаб s, кут θ, зсув):")
    b.formula(F_SIMILARITY)
    P("де x — точка поточного кадру, x′ — відповідна точка попереднього кадру, R(θ) — матриця повороту, "
      "t — вектор зсуву. RANSAC багаторазово оцінює модель за випадковими мінімальними підмножинами точок і "
      "залишає ту, з якою узгоджується найбільше точок, відкидаючи хибні відповідності. Добуток "
      "перетворень від кадру до кадру дає відображення T_{t} «кадр t → перший кадр», де A_{k} — перетворення між кадрами k і k − 1:")
    b.formula(F_CHAIN)
    P("У системі координат першого кадру нерухомі об'єкти залишаються на місці, а рухомі — зміщуються.")
    b.h2("2.4", "Розпізнавання номерних знаків")
    P("Загальна COCO-модель не має класу «номерний знак», тому використано YOLOv11s, донавчену на наборі "
      "номерних знаків (`morsetechlab/yolov11-license-plate-detection`, Hugging Face) [4]. Вона знаходить рамку "
      "знака, після чого символи читає EasyOCR [5] — двоетапна система, що складається з детектора тексту "
      "CRAFT [6] (оцінює для кожного пікселя ймовірність належності до символу та до проміжку між символами) "
      "і розпізнавача CRNN [7] (згорткові шари виділяють ознаки, двонаправлена LSTM моделює послідовність, "
      "декодування CTC перетворює її на рядок без потреби в сегментації символів).")
    P("Перед OCR знак збільшується втричі бікубічною інтерполяцією, переводиться у відтінки сірого та "
      "обробляється CLAHE [16] (адаптивним вирівнюванням гістограми з обмеженням контрасту), що підсилює "
      "контури дрібних символів. EasyOCR не знає формату номерів, тому застосовується постобробка за "
      "шаблоном українського номера `LLDDDDLL`, де L — одна з 12 літер, що мають однакове написання в "
      "кирилиці та латиниці (A, B, C, E, H, I, K, M, O, P, T, X), D — цифра.")
    b.h2("2.5", "Виявлення та розпізнавання облич. YuNet і SFace")
    P("YuNet [8] — мініатюрний згортковий детектор облич (≈ 75 тис. параметрів) з OpenCV Zoo [15], що для "
      "кожного обличчя повертає рамку, 5 ключових точок (центри очей, кінчик носа, кутики рота) та "
      "впевненість. За ключовими точками обличчя вирівнюється афінним перетворенням до стандартного "
      "положення розміром 112×112 пікселів.")
    P("SFace [9] — мережа розпізнавання облич, навчена з функцією втрат Sigmoid-Constrained Hypersphere Loss, "
      "яка перетворює вирівняне обличчя на 128-вимірний вектор ознак. Обличчя однієї людини дають близькі "
      "вектори, різних людей — далекі. Мірою близькості є косинусна схожість:")
    b.formula(F_COSINE)
    P("де a і b — 128-вимірні вектори ознак двох облич, ‖·‖ — евклідова норма.")
    P("Обличчя вважаються належними одній людині, якщо sim > 0,363 — поріг, рекомендований авторами моделі "
      "для косинусної метрики.")
    b.h2("2.6", "Обґрунтування вибору моделей")
    P(f"Підсумковий перелік використаних моделей і причини їх вибору наведено в таблиці {b.next_tab()}.")
    b.table("Використані моделі та методи",
            ["Задача", "Модель / метод", "Чому обрано"],
            [["Виявлення об'єктів", "YOLO11m, COCO (80 класів)",
              "сучасний точний одностадійний детектор; усі потрібні класи вже є в COCO"],
             ["Відстеження", "ByteTrack",
              "вбудований в Ultralytics; утримує треки при частковому перекритті"],
             ["Рух камери", "Ши–Томасі + Лукас–Канаде + RANSAC",
              "класичні швидкі методи OpenCV; модель подібності відповідає руху камери"],
             ["Номерний знак", "YOLOv11s-LP (morsetechlab)",
              "відома відкрита мережа, донавчена саме на номерних знаках"],
             ["Символи номера", "EasyOCR (CRAFT + CRNN)",
              "готовий OCR з латиницею та цифрами; не потребує донавчання"],
             ["Обличчя", "YuNet + SFace (OpenCV Zoo)",
              "легкі, точні, працюють у самому OpenCV без додаткових фреймворків"],
             ["Побудова сцени", "YOLO11m-seg",
              "маски екземплярів для вирізання людей і собаки з фото"]],
            [3.4, 4.6, 8.5])

    # ------------------------------------------------------------------ 3
    b.h1("3", "Програмна реалізація")
    b.h2("3.1", "Середовище та структура проєкту")
    P("Програмне забезпечення написане мовою Python 3.14 з використанням бібліотек PyTorch 2.14, "
      "Ultralytics 8.4, OpenCV 5.0 [14], EasyOCR 1.7, NumPy, pandas і matplotlib. Нейромережі YOLO виконуються "
      "на графічному процесорі Apple M1 через Metal Performance Shaders (MPS). Код організовано як пакет "
      f"`src/lab3`, кожен модуль якого відповідає за один крок конвеєра (таблиця {b.next_tab()}). Запуск — "
      "скрипт `run_lab3.py` (команди scene, process, all) або ноутбук `notebooks/Lab3_video_processing.ipynb`, "
      "який виконує кроки по черзі й показує проміжні результати.")
    b.table("Модулі програмного забезпечення",
            ["Модуль", "Призначення"],
            [["config.py", "шляхи, моделі, порогові значення в одному місці"],
             ["plate_render.py", "малювання номерного знака та перенесення його на фото гомографією"],
             ["scene.py", "генерація синтетичного відео та еталонної розмітки"],
             ["detection.py", "клас ObjectTracker: YOLO11m + ByteTrack"],
             ["camera_motion.py", "клас CameraMotionEstimator: оцінка руху камери"],
             ["plates.py", "клас PlateReader, функції normalize_plate і vote"],
             ["faces.py", "клас FaceIdentifier: YuNet + SFace"],
             ["analysis.py", "рішення по треках: рухомий об'єкт, водій, номер"],
             ["annotate.py", "побудова анотованого відео"],
             ["evaluation.py", "метрики якості за еталонною розміткою"],
             ["report_figures.py", "графіки для звіту"],
             ["pipeline.py", "запуск усього конвеєра, збереження результатів"]],
            [4.5, 12])
    b.h2("3.2", "Схема конвеєра обробки")
    P(f"Схему конвеєра наведено на рисунку {b.next_fig()}. Кожен кадр проходить через детектор об'єктів з "
      "трекером і оцінювач руху камери; кожен другий кадр (параметр `PROCESS_EVERY = 2`) додатково "
      "аналізується на номерний знак і обличчя. Після проходу по всьому відео результати агрегуються по треках.")
    b.figure(FIG / "fig_pipeline.png", "Схема конвеєра обробки відео", 16.5)
    P("Основний цикл обробки реалізовано у функції `process_video` (лістинг " + b.next_lst() + ").")
    b.listing("Фрагмент функції process_video — основний цикл (pipeline.py)",
              fragment(source_of("pipeline.py", "process_video"), "while True:", "f += 1"))

    b.h2("3.3", "Виявлення та відстеження об'єктів")
    P("Клас `ObjectTracker` (лістинг " + b.next_lst() + ") завантажує YOLO11m і залишає лише класи, що "
      "стосуються завдання. Виклик `model.track(..., persist=True)` зберігає стан ByteTrack між кадрами, тому "
      "кожен об'єкт отримує номер треку, що не змінюється. Результат кадру — список об'єктів `Detection` з "
      "номером треку, класом, впевненістю та рамкою. Мінімальна впевненість детекції — 0,35.")
    b.listing("Метод ObjectTracker.update (detection.py)",
              fragment(method_of("detection.py", "ObjectTracker", "update"), "def update", "for t, c, p, b"))

    b.h2("3.4", "Компенсація руху камери")
    P("Метод `CameraMotionEstimator.update` (лістинг " + b.next_lst() + ") реалізує алгоритм з підрозділу 2.3. "
      "Маска фону будується з рамок об'єктів, розширених на 8 пікселів, — точки на людях і тваринах не "
      "впливають на оцінку. Шукається до 500 кутових точок, оптичний потік обчислюється з вікном 21×21 на "
      "чотирьох рівнях піраміди, поріг RANSAC — 1,5 пікселя. Перетворення накопичуються в матрицю 3×3 "
      "«поточний кадр → перший кадр».")
    b.listing("Фрагмент методу CameraMotionEstimator.update (camera_motion.py)",
              fragment(method_of("camera_motion.py", "CameraMotionEstimator", "update"),
                       "if self.prev_gray is not None:", "self.to_ref = self.to_ref @ step"))

    b.h2("3.5", "Розпізнавання номерного знака")
    P("Клас `PlateReader` знаходить знак, готує зображення (розширення рамки на 4 % по ширині та 12 % по "
      "висоті, збільшення ×3, відтінки сірого, CLAHE) і передає його EasyOCR з алфавітом із великих "
      "латинських літер і цифр. Фрагменти тексту впорядковуються зліва направо й об'єднуються в сирий рядок.")
    P("Ключова частина — функція `normalize_plate` (лістинг " + b.next_lst() + "). Вона перебирає всі вікна "
      "по 8 символів сирого рядка і виправляє типові помилки OCR залежно від позиції: на місцях цифр O→0, "
      "I→1, B→8, S→5 тощо, на місцях літер — навпаки. Недопустима для українського номера літера штрафується. "
      "Обирається вікно з найменшою кількістю виправлень, а читання з більш ніж двома виправленнями "
      "відкидається як ненадійне. Так рядок «UA KA7429MB» (з кодом країни із синьої смуги) перетворюється на "
      "KA7429MB, а «7429MB» (половина знака перекрита пасажиром) відкидається.")
    b.listing("Нормалізація за шаблоном українського номера (plates.py)", source_of("plates.py", "normalize_plate"))
    P("Результати з різних кадрів об'єднуються функцією `vote` (лістинг " + b.next_lst() + ") двома способами: "
      "голосуванням за цілий рядок і посимвольним голосуванням, зваженим на впевненість OCR. Посимвольне "
      "голосування дає правильний номер навіть тоді, коли кожен кадр має одну помилку в різних позиціях.")
    b.listing("Фрагмент функції vote — посимвольне голосування (plates.py)",
              fragment(source_of("plates.py", "vote"), "valid = [r for r", 'by_char = "".join'))

    b.h2("3.6", "Ідентифікація водія")
    P("Клас `FaceIdentifier` (лістинг " + b.next_lst() + ") на етапі ініціалізації обчислює вектор SFace для "
      "найбільшого обличчя на еталонному фото водія. Для кожного кадру знаходяться всі обличчя з впевненістю "
      "YuNet ≥ 0,8, обличчя менші за 24 пікселі пропускаються, решта вирівнюється і порівнюється з еталоном. "
      "Кожне обличчя прив'язується до найменшої рамки `person`, у яку потрапляє його центр, — так схожість "
      "накопичується по треках людей.")
    b.listing("Метод FaceIdentifier.process (faces.py)",
              fragment(method_of("faces.py", "FaceIdentifier", "process"), "def process", "return out"))

    b.h2("3.7", "Прийняття рішень по треках")
    P("Функція `summarize` модуля `analysis.py` агрегує покадрові результати. Спочатку треки очищуються:")
    b.bullets(["**об'єднання фрагментів**: два треки одного класу, що майже не співіснують у часі (перекриття не "
               "більше 20 % коротшого) і мають однакові медіанні рамки в координатах першого кадру (IoU ≥ 0,6), "
               "вважаються одним об'єктом — це виправляє зміну номера треку під час наближення камери",
               "**відкидання коротких треків**: треки, коротші за 10 % відео, вважаються шумом",
               "**відкидання частин інших об'єктів**: трек, що у ≥ 80 % кадрів лежить усередині рамки довшого "
               "треку того ж класу",
               "**відкидання крайових треків**: середня впевненість < 0,5 і ≥ 80 % кадрів рамка впирається в "
               "край кадру (наприклад, піднята рука водія, що частково потрапила в кадр)"])
    P("**Рухомий об'єкт.** Рамки треку переводяться в систему координат першого кадру. Спостереження, де "
      "площа рамки менша за 70 % від 90-го перцентиля площ (об'єкт обрізаний краєм кадру або перекритий), "
      "відкидаються. Для кожної сторони рамки, що не торкається краю кадру, обчислюється розмах положень "
      "(різниця 90-го та 10-го перцентилів), береться медіана по сторонах і ділиться на характерний розмір "
      "об'єкта:")
    b.formula(F_MOVE)
    P("де e — сторона рамки (ліва, верхня, права, нижня), P_{90} і P_{10} — 90-й і 10-й перцентилі її "
      "положення, w і h — ширина і висота рамки в координатах першого кадру.")
    P("Об'єкт вважається рухомим, якщо M > 0,25. Медіана по сторонах робить оцінку стійкою до часткових "
      "перекриттів, оскільки перекриття змінює одну-дві сторони рамки, а не всі чотири. Для порівняння "
      "обчислюються також дві «наївні» оцінки без компенсації руху камери: зміщення центру рамки від першої "
      "до останньої появи (як у демонстрації) і максимальне зміщення центру від першої появи.")
    b.listing("Фрагмент функції movement_score (analysis.py)",
              fragment(source_of("analysis.py", "movement_score"), "refs = [_box_to_ref", "return float"))
    P("**Водій** — трек класу person, для якого є щонайменше 5 облич і середня схожість з еталоном "
      "найбільша; якщо вона не перевищує 0,363, водія не ідентифіковано. **Номер** приписується треку "
      "автомобіля, у рамці якого найчастіше знаходився центр знайденого знака.")

    b.h2("3.8", "Анотоване відео та оцінювання")
    P("Модуль `annotate.py` другим проходом по відео наносить на кадри рамки об'єктів з номерами треків, "
      "класами та впевненістю, позначки DRIVER (зелена рамка) і MOVING (червона рамка), рамку номерного "
      "знака з прочитаним після голосування номером, рамки облич зі значенням схожості та інформаційну "
      "панель. Траєкторія рухомого об'єкта зберігається в координатах першого кадру і в кожному кадрі "
      "переноситься оберненою матрицею руху камери, тому вона «прив'язана» до сцени, а не до екрана. Написи "
      "зроблено латиницею, оскільки шрифти Hershey в OpenCV не підтримують кирилицю. Модуль `evaluation.py` "
      "зіставляє треки з еталонними об'єктами за середнім IoU і обчислює метрики, наведені в розділі 4.")

    # ------------------------------------------------------------------ 4
    b.h1("4", "Результати")
    P(f"Обробка {META['frames']} кадрів на Apple M1 тривала {META['processing_s']:.0f} с "
      f"({f2(META['frames'] / META['processing_s'], 1)} кадр/с) з урахуванням OCR і розпізнавання облич "
      "у кожному другому кадрі. Побудова анотованого відео та ілюстрацій займає ще близько 1 хв.")
    b.h2("4.1", "Виявлені об'єкти")
    P(f"Трекер створив {S['n_tracks_total']} треків. Після очищення об'єднано фрагменти "
      f"{', '.join(f'#{k}→#{v}' for k, v in S['merged_tracks'].items()) or '—'}, відкинуто частини інших "
      f"об'єктів {', '.join('#' + str(t) for t in S['part_tracks']) or '—'} та крайові треки "
      f"{', '.join('#' + str(t) for t in S['edge_tracks']) or '—'}, решта коротких треків відкинута як шум. "
      f"Залишилося {len(TRACKS)} об'єктів, що точно відповідають вмісту сцени (таблиця {b.next_tab()}).")
    role = {"ВОДІЙ": "водій", "інша людина (пасажир)": "пасажир", "інший автомобіль": "другий автомобіль",
            "додатковий об'єкт": "додатковий об'єкт (собака)"}
    b.table("Підсумкові треки об'єктів",
            ["Трек", "Клас", "Кадрів", "Впевн.", "Оцінка руху M", "Рух", "Схожість з водієм (облич)",
             "Роль"],
            [[f"#{t['track_id']}", t["cls"], t["n_frames"], f2(t["mean_conf"]), f2(t["move_score"]),
              "так" if t["moving"] else "ні",
              "—" if t["mean_similarity"] is None else f"{f2(t['mean_similarity'], 3)} ({t['face_samples']})",
              role.get(t["role"], t["role"])] for t in TRACKS],
            [1.5, 1.6, 1.6, 1.8, 1.9, 1.1, 2.9, 4.1], size=11,
            align_cols={0: 1, 2: 1, 3: 1, 4: 1, 5: 1, 6: 1})
    P(f"Впевненість детектора по кадрах показано на рисунку {b.next_fig()}. Для автомобіля з номером, водія та "
      "собаки вона стабільно перевищує 0,9. Помітні провали на 11–13 с відповідають моменту, коли пасажир "
      "проходить перед собакою і передньою частиною автомобіля, а провал пасажира на 24–26 с — моменту, коли "
      "він заходить за водія. Другий автомобіль видно лише частково (він обрізаний лівим краєм кадру), тому "
      "його впевненість нижча — близько 0,8.")
    b.figure(FIG / "fig_confidence.png", "Впевненість YOLO11m по кадрах для кожного треку", 16.5)
    P(f"Приклади кадрів анотованого відео наведено на рисунку {b.next_fig()}.")
    _annot_grid()
    b.figure(FIG / "annot_grid.jpg", "Кадри анотованого відео (2,4 с; 12 с; 20,8 с; 28 с)", 16.5)

    b.h2("4.2", "Рухомий об'єкт")
    P(f"Оцінки руху для всіх треків трьома способами наведено на рисунку {b.next_fig()}. З компенсацією руху "
      f"камери пасажир має оцінку M = {f2(MOV['move_score'])}, а всі нерухомі об'єкти — не більше "
      f"{f2(max(t['move_score'] for t in TRACKS if not t['moving']))}, тобто рішення має великий запас "
      "відносно порогу 0,25.")
    b.figure(FIG / "fig_movement.png", "Оцінки руху об'єктів з компенсацією і без компенсації руху камери", 16.5)
    naive2 = [t for t in TRACKS if t["move_score_naive_max"] > 0.25 and not t["moving"]]
    P("Без компенсації картина інша. Оцінка «максимальне зміщення» помилково вважає рухомими "
      + ", ".join(f"#{t['track_id']} ({role.get(t['role'], t['role'])}, {f2(t['move_score_naive_max'])})"
                  for t in naive2)
      + " — під час наближення камери ці об'єкти зміщуються в кадрі більше, ніж на чверть власного розміру. "
      "Оцінка з демонстрації (перша → остання поява) у нашому відео дала правильну відповідь лише тому, що "
      "камера в кінці повертається у вихідне положення; при іншій траєкторії камери вона б також помилилася. "
      "Це підтверджує необхідність компенсації руху камери, про яку згадано в рекомендаціях демонстрації.")
    P(f"Траєкторії центрів усіх треків у координатах першого кадру показано на рисунку {b.next_fig()}: "
      "нерухомі об'єкти утворюють компактні «плями», а пасажир — довгу горизонтальну лінію. Злам наприкінці "
      "траєкторії відповідає моменту, коли пасажир заходить за водія і видно лише його голову.")
    b.figure(FIG / "fig_trajectory.png", "Траєкторії об'єктів у системі координат першого кадру", 15.5)
    cam = EV["camera"]
    P(f"Точність оцінки руху камери перевірено за еталоном (рисунок {b.next_fig()}). Оцінене наближення "
      f"практично збігається зі справжнім. Середня похибка перенесення точок кадру в систему першого кадру "
      f"становить {f2(cam['mean_error_px'])} пікселя, максимальна — {f2(cam['max_error_px'])} пікселя "
      "наприкінці відео. Похибка повільно накопичується через добуток перетворень, але залишається на два "
      "порядки меншою за розміри об'єктів.")
    b.figure(FIG / "fig_camera.png", "Оцінене і справжнє наближення камери та похибка оцінки", 15.5)

    b.h2("4.3", "Номер автомобіля")
    pl, plv = EV["plate"], S["plate"]
    counts = ", ".join(f"{k}: {v}" for k, v in plv["counts"].items())
    P(f"Детектор номерних знаків знайшов знак у {plv['frames_with_plate']} проаналізованих кадрах, з них "
      f"{plv['valid_reads']} читань відповідали шаблону українського номера. Результати голосування: {counts}. "
      f"Голосування за рядком дало {plv['by_string']}, посимвольне — {plv['by_char']}. "
      f"**Номер автомобіля: {S['plate_text']}**, він приписаний треку #{S['plate_car_id']} "
      "(автомобіль ЗАЗ Sens), що відповідає еталону.")
    P(f"Приклади підготовки знака до OCR і результати читання наведено на рисунку {b.next_fig()}. У кадрі 300 "
      "пасажир перекриває ліву половину знака: OCR повертає «7429MB», і нормалізація правильно відкидає "
      "це читання замість того, щоб видати хибний номер. Код «UA» із синьої смуги, який OCR інколи читає "
      "разом із номером, відкидається пошуком вікна з 8 символів.")
    b.figure(FIG / "fig_plate_examples.png", "Знайдений номерний знак, підготовлене зображення та результат OCR", 16)

    b.h2("4.4", "Ідентифікація водія")
    P(f"Схожість облич з еталоном у часі показано на рисунку {b.next_fig()}. Для водія (трек "
      f"#{DRV['track_id']}) вона стабільно тримається в межах "
      f"{f2(EV['driver']['similarity']['driver']['min'], 3)}–{f2(EV['driver']['similarity']['driver']['max'], 3)}, "
      f"середнє {f2(DRV['mean_similarity'], 3)} у {DRV['face_samples']} кадрах, тобто значно вище порогу 0,363. "
      f"Для пасажира середня схожість {f2(PAS_SIM['mean'], 3)} (максимум {f2(PAS_SIM.get('max', 0), 3)}) — "
      "значно нижче порогу. **Водія ідентифіковано: трек "
      f"#{S['driver_id']}, середня схожість {f2(S['driver_similarity'], 3)} > 0,363.**")
    b.figure(FIG / "fig_similarity.png", "Косинусна схожість облич з еталонним фото водія в часі", 16)
    P(f"Вирівняні обличчя, з якими працює SFace, наведено на рисунку {b.next_fig()}. Хоча еталонне фото "
      "зроблене через чотири роки, в іншому одязі та за іншого освітлення, схожість облич водія з еталоном "
      "≈ 0,7. Обличчя пасажира частково закрите сонцезахисними окулярами, і його схожість близька до нуля.")
    b.figure(FIG / "fig_faces.png", "Еталонне обличчя та вирівняні обличчя з відео", 16.5)

    b.h2("4.5", "Кількісні показники")
    P(f"Метрики якості, обчислені за еталонною розміткою, наведено в таблицях {b.next_tab()} і "
      f"{b.next_tab()[:-1]}{int(b.next_tab()[-1]) + 1}. Об'єкт вважається знайденим у кадрі, якщо є детекція "
      "відповідного класу з IoU ≥ 0,5 з еталонною рамкою; враховуються кадри з видимістю об'єкта ≥ 30 %.")
    names = {"car_main": "автомобіль з номером", "car_second": "другий автомобіль", "driver": "водій",
             "passenger": "пасажир", "dog": "собака"}
    b.table("Якість виявлення та відстеження об'єктів",
            ["Об'єкт", "Кадрів у полі зору", "Частка знайдених (IoU ≥ 0,5)", "Середній IoU", "Номери треків"],
            [[names[k], v["frames_visible"], f2(v["recall_iou50"], 3), f2(v["mean_iou"], 3),
              ", ".join(f"#{tid} ({n})" for tid, n in v["track_ids"].items())]
             for k, v in EV["detection"].items()],
            [4, 2.6, 3.2, 2.4, 4.3], size=11, align_cols={1: 1, 2: 1, 3: 1})
    b.table("Підсумкові показники за пунктами завдання",
            ["Показник", "Значення"],
            [["Кадрів з відкритим знаком (видимість ≥ 90 %), проаналізованих на номер",
              pl["frames_analysed_plate_visible"]],
             ["Частка кадрів, де знак знайдено", f"{f2(pl['detection_rate'] * 100, 1)} %"],
             ["Частка кадрів з читанням правильного формату", f"{f2(pl['valid_format_rate'] * 100, 1)} %"],
             ["Частка кадрів, де номер прочитано точно", f"{f2(pl['exact_read_rate'] * 100, 1)} %"],
             ["Підсумковий номер (голосування) / еталон", f"{S['plate_text']} / {pl['truth']}"],
             ["Ідентифікований водій відповідає еталону", "так" if EV["driver"]["correct"] else "ні"],
             ["Середня схожість: водій / пасажир", f"{f2(DRV['mean_similarity'], 3)} / {f2(PAS_SIM['mean'], 3)}"],
             ["Рухомий об'єкт (з компенсацією) відповідає еталону", "так" if EV["moving"]["correct"] else "ні"],
             ["Те саме без компенсації (максимальне зміщення)",
              "так" if EV["moving"]["naive_max_correct"] else
              "ні — рухомими визнано " + ", ".join(names.get(x, x) for x in EV["moving"]["naive_max_predicted"])],
             ["Похибка оцінки руху камери: середня / максимальна",
              f"{f2(cam['mean_error_px'])} / {f2(cam['max_error_px'])} px"],
             ["Швидкість обробки (Apple M1)", f"{f2(META['frames'] / META['processing_s'], 1)} кадр/с"]],
            [10.5, 6], size=11)
    drv_det = EV["detection"]["driver"]["recall_iou50"]
    P(f"Нижча частка знайдених кадрів для водія ({f2(drv_det, 3)}) пояснюється тим, що під час наближення "
      "камери водія видно лише частково: еталонна рамка видимої частини включає підняту руку біля краю кадру, "
      "а детектор її не охоплює, тож IoU опускається нижче 0,5. Сам трек водія при цьому не переривається "
      "(трек #2 існує у всіх 1000 кадрах). Пасажир розбитий на два треки: після проходу за водієм частина "
      "його фігури знову з'являється і отримує новий номер, який відкидається як короткий.")

    # ------------------------------------------------------------------ 5
    b.h1("5", "Аналіз помилок та способи покращення")
    P("Незважаючи на правильні підсумкові рішення, на рівні окремих кадрів виявлено такі помилки та обмеження:")
    b.numbered([
        "**Зміна номера треку (ID switch).** Під час швидкого наближення камери ByteTrack присвоїв головному "
        "автомобілю новий номер. Виправлено об'єднанням фрагментів у координатах першого кадру; у загальному "
        "випадку допоможе трекер з компенсацією руху камери (BoT-SORT з GMC) або ознаками зовнішнього вигляду "
        "(re-identification)",
        "**Хибні та часткові детекції.** Піднята рука водія біля краю кадру детектувалася як окрема людина з "
        "низькою впевненістю. Відфільтровано правилом для крайових треків",
        "**Хибне обличчя.** YuNet в окремих кадрах знаходив «обличчя» на морді собаки (схожість ≈ 0,1). Воно "
        "не впливає на результат, бо не потрапляє в рамку жодної людини; можна додатково підвищити поріг "
        "впевненості YuNet або перевіряти клас об'єкта, якому належить обличчя",
        "**Перекриття номера.** Коли пасажир або собака перекривають знак, OCR читає лише частину символів. "
        "Шаблон номера і голосування по кадрах повністю усувають вплив таких кадрів. EasyOCR не знає формату "
        "номерів, тому постобробка обов'язкова; кращий варіант — спеціалізований OCR для номерних знаків "
        "(наприклад, fast-plate-ocr) або донавчання CRNN на українських номерах",
        "**Накопичення похибки руху камери.** Похибка зростає приблизно на 0,1 пікселя за секунду; для довших "
        "відео варто періодично зіставляти кадр безпосередньо з опорним (ключові кадри) або використовувати "
        "гомографію замість моделі подібності при сильних змінах ракурсу",
        "**Окуляри та ракурс обличчя.** Сонцезахисні окуляри пасажира знижують інформативність обличчя. Для "
        "водія з окулярами або в профіль схожість може опуститися нижче порогу в окремих кадрах, тому рішення "
        "приймається за середнім значенням по треку",
        "**Синтетичність сцени.** Вирізані об'єкти не мають власного руху кінцівок і реалістичних тіней, а "
        "освітлення вирівнюється лише статистично. На реальному відео слід очікувати нижчої впевненості "
        "детекторів, розмиття в русі та відблисків на номері; рекомендації щодо зйомки — номер ширший за "
        "150 пікселів, обличчя водія анфас, камера відносно нерухома",
        "**Ліцензії.** Модель номерних знаків і YOLO11 поширюються за ліцензією AGPL-3.0: для навчання "
        "використання дозволене, для комерційного продукту потрібна інша ліцензія"])

    # ------------------------------------------------------------------ ВИСНОВКИ
    b.structural("Висновки")
    P("У роботі розроблено програмне забезпечення для обробки відео, яке з використанням відомих попередньо "
      "навчених нейронних мереж виконує всі пункти завдання. Для перевірки сформовано синтетичний відеофільм "
      "тривалістю 40 с (1000 кадрів, 1280×720), що відрізняється від демонстраційного прикладу автомобілем, "
      "номером, водієм, додатковими об'єктами, тривалістю та рухом камери, і для якого є точна еталонна "
      "розмітка.")
    P(f"Мережа YOLO11m разом з трекером ByteTrack виділила й розпізнала всі {len(TRACKS)} об'єктів сцени: "
      "автомобіль з номером, другий автомобіль, водія, пасажира та собаку; частка кадрів зі знайденим "
      f"об'єктом становить {f2(min(v['recall_iou50'] for v in EV['detection'].values()), 2)}–"
      f"{f2(max(v['recall_iou50'] for v in EV['detection'].values()), 2)} при середньому IoU "
      f"{f2(min(v['mean_iou'] for v in EV['detection'].values()), 2)}–"
      f"{f2(max(v['mean_iou'] for v in EV['detection'].values()), 2)}.")
    P(f"Мережа YOLOv11-LP знайшла номерний знак у {f2(pl['detection_rate'] * 100, 0)} % кадрів, де він відкритий, "
      f"а EasyOCR з постобробкою за шаблоном українського номера прочитав його точно в "
      f"{f2(pl['exact_read_rate'] * 100, 0)} % таких кадрів. Голосування по кадрах дало номер "
      f"**{S['plate_text']}**, що збігається з еталоном.")
    P(f"Мережі YuNet і SFace за окремим еталонним фото ідентифікували водія (трек #{S['driver_id']}): середня "
      f"косинусна схожість {f2(DRV['mean_similarity'], 3)} проти {f2(PAS_SIM['mean'], 3)} для пасажира при "
      "порозі 0,363.")
    P(f"Єдиний рухомий об'єкт — пасажир — визначено з оцінкою {f2(MOV['move_score'])} при порозі 0,25. "
      "Показано, що без компенсації руху камери нерухомі об'єкти (собака, другий автомобіль) помилково "
      "визнаються рухомими; запропонований метод на основі оптичного потоку та RANSAC оцінює рух камери з "
      f"середньою похибкою {f2(cam['mean_error_px'])} пікселя.")
    P("Головні інженерні рішення, що забезпечили правильний результат: прийняття рішень по треку, а не по "
      "окремому кадру; постобробка OCR за шаблоном номера з голосуванням; компенсація руху камери; очищення "
      "треків (об'єднання фрагментів, відкидання хибних і часткових детекцій). Напрями подальшого "
      "покращення — спеціалізований OCR для номерів, трекер з компенсацією руху камери та ознаками "
      "зовнішнього вигляду, перевірка на реальному відео.")

    # ------------------------------------------------------------------ ДЖЕРЕЛА
    b.structural("Перелік джерел посилання")
    refs = REFERENCES
    for i, r in enumerate(refs, 1):
        P(f"{i}. {r}")

    # ------------------------------------------------------------------ ДОДАТКИ
    b.chapter = "А"
    p = b.structural("Додаток А")
    b.toc[-1] = (0, "ДОДАТОК А Вихідний код основних модулів")
    t = b.para("**Вихідний код основних модулів**", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    P("Повний код проєкту (включно з генератором сцени `scene.py`, рендером номера `plate_render.py`, "
      "модулями `annotate.py`, `evaluation.py`, `report_figures.py` і ноутбуком) знаходиться в репозиторії "
      "роботи. Нижче наведено модулі, що реалізують обробку відео.")
    for mod in ["config.py", "detection.py", "camera_motion.py", "plates.py", "faces.py", "analysis.py",
                "pipeline.py"]:
        b.listing(f"Модуль {mod}", (SRC / mod).read_text(), size=8.5)

    b.chapter = "Б"
    b.structural("Додаток Б")
    b.toc[-1] = (0, "ДОДАТОК Б Тези доповіді")
    b.para("**Тези доповіді про виконану роботу**", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    b.numbered([
        "**Завдання.** Обробити відео 30–60 с: виділити й розпізнати об'єкти, прочитати номер автомобіля, "
        "ідентифікувати водія, визначити рухомий об'єкт",
        "**Відео.** Синтетична 40-секундна сцена: ЗАЗ Sens з номером KA 7429 MB, другий автомобіль, водій, "
        "пасажир, що йде (рухомий об'єкт), собака; камера «з руки» з наближенням; є еталонна розмітка",
        "**Конвеєр.** YOLO11m + ByteTrack → об'єкти й треки; оптичний потік + RANSAC → рух камери; "
        "YOLOv11-LP + EasyOCR + шаблон + голосування → номер; YuNet + SFace → водій",
        f"**Результати.** {len(TRACKS)} об'єктів знайдено правильно; номер {S['plate_text']} прочитано точно в "
        f"{f2(pl['exact_read_rate'] * 100, 0)} % кадрів; водій — схожість {f2(DRV['mean_similarity'], 2)} проти "
        f"{f2(PAS_SIM['mean'], 2)}; рухомий — пасажир (оцінка {f2(MOV['move_score'])} при порозі 0,25)",
        "**Головний висновок.** Без компенсації руху камери нерухомі об'єкти визнаються рухомими; рішення по "
        "треку та постобробка OCR роблять систему стійкою до помилок окремих кадрів",
        "**Демонстрація.** Анотоване відео data/synthetic_scene_annotated.mp4 з рамками, позначками DRIVER "
        "і MOVING, номером і траєкторією рухомого об'єкта"])
    return b


# ==========================================================================
# Композиції кадрів для звіту
# ==========================================================================
def _grid(paths, out: Path, cols=2, w=800):
    import cv2
    import numpy as np
    ims = [cv2.resize(cv2.imread(str(p)), (w, w * 9 // 16)) for p in paths]
    pad = np.full((w * 9 // 16, 8, 3), 255, np.uint8)
    rows = []
    for i in range(0, len(ims), cols):
        row = []
        for j, im in enumerate(ims[i:i + cols]):
            row += ([pad] if j else []) + [im]
        rows.append(np.hstack(row))
    hpad = np.full((8, rows[0].shape[1], 3), 255, np.uint8)
    out_img = rows[0]
    for r in rows[1:]:
        out_img = np.vstack([out_img, hpad, r])
    cv2.imwrite(str(out), out_img, [cv2.IMWRITE_JPEG_QUALITY, 90])


def _frames_grid():
    _grid([FIG / f"scene_frame_{f:04d}.jpg" for f in (0, 250, 450, 900)], FIG / "scene_grid.jpg")


def _annot_grid():
    _grid([FIG / f"annotated_{f:04d}.jpg" for f in (60, 300, 520, 700)], FIG / "annot_grid.jpg")


# ==========================================================================
# Два проходи: визначення сторінок для змісту
# ==========================================================================
def _pdf_pages(docx: Path) -> list[str]:
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx)],
                       check=True, capture_output=True)
        pdf = Path(tmp) / (docx.stem + ".pdf")
        txt = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True,
                             text=True).stdout
        shutil.copy(pdf, OUT.with_suffix(".pdf"))
    return txt.split("\f")


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def find_pages(b: Builder, pages: list[str]) -> dict[str, int]:
    """Шукає кожен заголовок змісту в тексті сторінок (після сторінки змісту)."""
    norm_pages = [_norm(p) for p in pages]
    # Пошук починаємо після сторінок змісту (сторінки з рядками-заповнювачами «.....»)
    toc_pages = [i for i, p in enumerate(pages) if len(re.findall(r"\.{10,}\s*\d+", p)) >= 3]
    start = max(toc_pages) + 1
    result = {}
    for _, text in b.toc:
        key = _norm(text)
        probe = key.split(" Вихідний")[0].split(" Тези")[0]
        for i in range(start, len(norm_pages)):
            if probe in norm_pages[i]:
                result[text] = i + 1
                start = i
                break
    return result


def _stats(b: Builder, n_pages: int) -> dict:
    return {"pages": n_pages,
            "figures": sum(v for k, v in b.fig_n.items()),
            "tables": sum(v for k, v in b.tab_n.items() if not k.startswith("_formula_"))}


def main():
    """Збирає звіт, доки номери сторінок у змісті та рефераті не стабілізуються (≤ 4 проходи)."""
    pages, stats = None, None
    for _ in range(4):
        b = build(pages, stats)
        b.save(OUT)
        text_pages = _pdf_pages(OUT)
        n_pages = len([p for p in text_pages if p.strip()])
        new_pages, new_stats = find_pages(b, text_pages), _stats(b, n_pages)
        if new_pages == pages and new_stats == stats:
            break
        pages, stats = new_pages, new_stats
    else:
        print("Увага: номери сторінок не стабілізувались")
    print(f"Звіт збережено: {OUT} ({stats['pages']} с., {stats['figures']} рис., {stats['tables']} табл.)")
    for k, v in pages.items():
        print(f"  {v:>3}  {k}")


if __name__ == "__main__":
    sys.exit(main())
